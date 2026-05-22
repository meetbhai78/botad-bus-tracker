# -*- coding: utf-8 -*-
"""Generate Botad Bus Tracker full project documentation as Word (.docx)."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date

def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p

def add_bullet(doc, text):
    doc.add_paragraph(text, style='List Bullet')

def add_numbered(doc, text):
    doc.add_paragraph(text, style='List Number')

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()

def build():
    doc = Document()
    title = doc.add_heading('Botad Bus Tracker — Complete Project Documentation', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, f'Version 1.0 | Date: {date.today().strftime("%d %B %Y")} | City: Botad, Gujarat, India')
    add_para(doc, 'Yeh document poora project kaise kaam karta hai — har chhoti detail ke saath — explain karta hai.')

    # TOC placeholder
    add_heading(doc, 'Table of Contents', 1)
    toc_items = [
        '1. Project Overview', '2. System Architecture', '3. Technology Stack',
        '4. Folder Structure', '5. Database (MongoDB) Schemas', '6. Authentication & JWT',
        '7. Backend REST API (Har Endpoint)', '8. Socket.io Real-Time Events',
        '9. AI ETA Microservice', '10. Seed Data', '11. React Web Frontend',
        '12. Flutter Driver App', '13. Flutter Passenger App', '14. Complete User Flows',
        '15. Environment Variables', '16. Setup & Run (Step by Step)',
        '17. Deployment', '18. Security', '19. Known Issues & Tips', '20. Glossary',
    ]
    for item in toc_items:
        add_bullet(doc, item)

    # 1 OVERVIEW
    add_heading(doc, '1. Project Overview', 1)
    add_para(doc, 'Botad Bus Tracker ek real-time city bus tracking system hai jo Botad (Gujarat) ke liye banaya gaya hai. Isme teen roles hain: Admin (web), Driver (mobile app), Passenger (mobile app + web map).')
    add_para(doc, 'Main features:', bold=True)
    features = [
        'Live GPS tracking — driver phone se location har 3 second par server par jati hai',
        'Passengers live map par bus dekh sakte hain',
        'Admin stops, routes, timetable, buses, drivers manage karta hai',
        'Digital ticket booking + QR code',
        'Driver QR scan karke ticket verify karta hai',
        'ETA (Estimated Time of Arrival) — math + optional AI service',
        'Timetable search — kis route par kaunse stop par bus kab aayegi',
    ]
    for f in features:
        add_bullet(doc, f)

    add_heading(doc, '1.1 Project Components', 2)
    add_table(doc, ['Component', 'Folder', 'Technology', 'Port'],
        [
            ['Backend API', 'backend/', 'Node.js, Express, Socket.io', '5000'],
            ['Web Frontend', 'frontend/', 'React 18, Vite, Tailwind', '3000'],
            ['Driver App', 'driver-app/', 'Flutter', 'Mobile'],
            ['Passenger App', 'passenger-app/', 'Flutter', 'Mobile'],
            ['AI ETA Service', 'ai-eta/', 'Python FastAPI', '8000'],
            ['Database', 'MongoDB Atlas / Local', 'MongoDB', '27017'],
        ])

    # 2 ARCHITECTURE
    add_heading(doc, '2. System Architecture', 1)
    add_para(doc, 'Data flow (high level):')
    add_numbered(doc, 'Driver app GPS ON → har 3 sec Socket event driver:location → Backend MongoDB update → sab clients ko buses:locations broadcast')
    add_numbered(doc, 'Passenger / Web map Socket se bus markers update karte hain')
    add_numbered(doc, 'Ticket book → MongoDB Ticket + QR JSON → Driver scan → verify API → status used')
    add_numbered(doc, 'Admin web se CRUD → REST API → MongoDB')

    add_heading(doc, '2.1 Architecture Diagram (Text)', 2)
    add_para(doc, '''
┌─────────────┐  ┌─────────────┐  ┌──────────────┐
│ Driver App  │  │ Passenger   │  │ React Web    │
│  (Flutter)  │  │ App Flutter │  │ Admin+Map    │
└──────┬──────┘  └──────┬──────┘  └──────┬───────┘
       │ REST + Socket   │ REST + Socket   │ REST + Socket
       └─────────────────┼─────────────────┘
                         ▼
              ┌──────────────────────┐
              │  Express + Socket.io │
              │      (Port 5000)     │
              └──────────┬───────────┘
                         │
         ┌───────────────┼───────────────┐
         ▼               ▼               ▼
    ┌─────────┐   ┌───────────┐   ┌──────────┐
    │ MongoDB │   │ ai-eta    │   │ OSRM     │
    │         │   │ (optional)│   │ (routes) │
    └─────────┘   └───────────┘   └──────────┘
''')

    # 3 STACK
    add_heading(doc, '3. Technology Stack', 1)
    add_table(doc, ['Layer', 'Technologies'],
        [
            ['Backend', 'Node.js, Express 4, Mongoose 8, Socket.io 4, JWT, bcrypt, QRCode'],
            ['Database', 'MongoDB'],
            ['Web', 'React 18, Vite, Tailwind CSS, Axios, Recharts, Google Maps API'],
            ['Mobile', 'Flutter, flutter_map, latlong2, socket_io_client, mobile_scanner, qr_flutter, geolocator'],
            ['ETA AI', 'Python 3, FastAPI, Uvicorn'],
            ['CI/CD', 'GitHub Actions (Flutter APK build)'],
            ['Deploy', 'Render (API), Vercel (web), MongoDB Atlas'],
        ])

    # 4 FOLDER
    add_heading(doc, '4. Complete Folder Structure', 1)
    add_para(doc, '''botad-bus-tracker/
├── README.md, SETUP-WINDOWS.md
├── .github/workflows/build_apps.yml
├── ai-eta/ (main.py, requirements.txt)
├── backend/
│   ├── server.js (main entry)
│   ├── package.json
│   ├── .env.example
│   ├── public/legacy_admin.js (deprecated)
│   ├── tests/app.test.js
│   └── src/
│       ├── config/db.js
│       ├── controllers/ (auth, bus, trip, ticket, admin, route, routeAdmin, stop, timetable)
│       ├── middleware/ (auth.js, errorHandler.js)
│       ├── models/ (User, Bus, Route, Stop, Trip, Ticket, Timetable)
│       ├── routes/ (auth, buses, routes, trips, tickets, stops, timetable, admin)
│       ├── services/etaService.js
│       ├── socket/index.js
│       └── utils/ (seed.js, distance.js, generateToken.js)
├── frontend/src/ (App.jsx, pages/, admin/, hooks/, services/)
├── driver-app/lib/ (screens/, services/, constants.dart)
└── passenger-app/lib/ (screens/, services/, widgets/, theme/)''')

    # 5 DATABASE
    add_heading(doc, '5. Database Schemas (MongoDB)', 1)

    add_heading(doc, '5.1 User Collection', 2)
    add_table(doc, ['Field', 'Type', 'Description'],
        [
            ['name', 'String', 'User ka naam'],
            ['phone', 'String (unique)', 'Login ID — phone number'],
            ['email', 'String', 'Optional email'],
            ['password', 'String', 'bcrypt hash (12 rounds)'],
            ['role', 'enum', 'passenger | driver | admin'],
            ['fcmToken', 'String', 'Push notification (future)'],
            ['profilePhoto', 'String', 'Photo URL'],
            ['createdAt', 'Date', 'Auto timestamp'],
        ])

    add_heading(doc, '5.2 Bus Collection', 2)
    add_table(doc, ['Field', 'Type', 'Description'],
        [
            ['busNumber', 'String', 'Number plate e.g. GJ-11-T-2847'],
            ['busName', 'String', 'Short name e.g. B1'],
            ['capacity', 'Number', 'Total seats (default 50)'],
            ['driver', 'ObjectId → User', 'Assigned driver'],
            ['route', 'ObjectId → Route', 'Current route'],
            ['currentLocation', 'Object', 'lat, lng, speed, heading, updatedAt'],
            ['status', 'enum', 'active | idle | maintenance'],
            ['isLive', 'Boolean', 'Live GPS broadcasting'],
        ])

    add_heading(doc, '5.3 Route Collection', 2)
    add_table(doc, ['Field', 'Type', 'Description'],
        [
            ['routeNumber', 'String', 'R1, R2, R3'],
            ['routeName', 'String', 'Human readable name'],
            ['stops[]', 'Array', 'Embedded: name, lat, lng, order, estimatedTime'],
            ['totalDistance', 'Number', 'km'],
            ['totalTime', 'Number', 'minutes'],
            ['polyline', 'String', 'OSRM encoded polyline for map line'],
        ])

    add_heading(doc, '5.4 Stop Collection', 2)
    add_table(doc, ['Field', 'Type', 'Description'],
        [['name', 'String', 'Stop name'], ['lat', 'Number', 'Latitude'], ['lng', 'Number', 'Longitude'],
         ['address', 'String', 'Optional'], ['isActive', 'Boolean', 'Soft delete'], ['createdAt', 'Date', 'Auto']])

    add_heading(doc, '5.5 Trip Collection', 2)
    add_table(doc, ['Field', 'Type', 'Description'],
        [
            ['bus, driver, route', 'ObjectId', 'References'],
            ['startTime, endTime', 'Date', 'Trip duration'],
            ['passengersCount', 'Number', 'Booked passengers on this trip'],
            ['locationHistory[]', 'Array', 'GPS points during trip'],
            ['status', 'enum', 'scheduled | ongoing | completed'],
        ])

    add_heading(doc, '5.6 Ticket Collection', 2)
    add_table(doc, ['Field', 'Type', 'Description'],
        [
            ['passenger', 'ObjectId', 'Who booked'],
            ['trip', 'ObjectId', 'Optional — book-simple mein null'],
            ['fromStop, toStop', 'String', 'Stop names'],
            ['price', 'Number', 'Rupees (default ₹15)'],
            ['qrCode', 'String', 'JSON string ya QR image data URL'],
            ['status', 'enum', 'active | used | cancelled'],
            ['bookedAt', 'Date', 'Booking time'],
            ['expiresAt', 'Date', '24 hours validity'],
        ])

    add_heading(doc, '5.7 Timetable Collection', 2)
    add_table(doc, ['Field', 'Type', 'Description'],
        [
            ['route, bus', 'ObjectId', 'Which route/bus'],
            ['label', 'String', 'e.g. Morning 07:00'],
            ['departureTime', 'String', 'HH:mm'],
            ['direction', 'String', 'Description'],
            ['schedule[]', 'Array', 'stop, stopName, arrivalTime, order'],
            ['isActive', 'Boolean', 'Active flag'],
        ])

    # 6 AUTH
    add_heading(doc, '6. Authentication & JWT', 1)
    add_para(doc, 'Login/Register par server JWT token deta hai. Har protected API par header: Authorization: Bearer <token>')
    add_heading(doc, '6.1 Token Generation', 2)
    add_para(doc, 'jwt.sign({ id: userId }, JWT_SECRET, { expiresIn: "7d" }) — file: backend/src/utils/generateToken.js')
    add_heading(doc, '6.2 Middleware protect', 2)
    add_numbered(doc, 'Header se Bearer token nikalo')
    add_numbered(doc, 'jwt.verify se decode karo')
    add_numbered(doc, 'User DB se load karo (password exclude)')
    add_numbered(doc, 'req.user par set karo')
    add_heading(doc, '6.3 authorize(...roles)', 2)
    add_para(doc, 'Example: authorize("driver") — sirf driver role allowed')
    add_heading(doc, '6.4 Rate Limiting', 2)
    add_para(doc, 'Auth routes par 5 requests per minute per IP (express-rate-limit)')
    add_heading(doc, '6.5 Demo Logins (after npm run seed)', 2)
    add_table(doc, ['Role', 'Phone', 'Password', 'Where to use'],
        [
            ['Admin', '9999999999', 'admin123', 'Web http://localhost:3000/login'],
            ['Driver', '9876543210', 'driver123', 'Driver Flutter app'],
            ['Passenger', '9123456789', 'pass123', 'Passenger Flutter app'],
        ])

    # 7 API
    add_heading(doc, '7. Backend REST API — Complete List', 1)
    add_para(doc, 'Base URL: http://localhost:5000/api (production: Render URL)')

    add_heading(doc, '7.1 Health', 2)
    add_para(doc, 'GET /health — No auth. Response: { success: true, service: "Botad Bus Tracker API", city: "Botad, Gujarat" }')

    add_heading(doc, '7.2 Auth — /api/auth', 2)
    add_table(doc, ['Method', 'Path', 'Auth', 'Body', 'Response'],
        [
            ['POST', '/register', 'Public', 'name, phone, password, email?', 'token + user (passenger only)'],
            ['POST', '/login', 'Public', 'phone, password', 'token + user (any role)'],
            ['POST', '/driver/login', 'Public', 'phone, password', 'token + user (driver only)'],
            ['GET', '/me', 'Bearer', '-', 'current user'],
            ['POST', '/refresh-token', 'Bearer', '-', 'new token'],
        ])

    add_heading(doc, '7.3 Buses — /api/buses', 2)
    add_table(doc, ['Method', 'Path', 'Auth', 'Notes'],
        [
            ['GET', '/', 'Public', 'All buses. Query ?from=&to= filters by route stop order'],
            ['GET', '/nearby', 'Public', '?lat=&lng= live buses with distanceKm, etaMinutes'],
            ['GET', '/:id', 'Public', 'Single bus with driver + route'],
            ['GET', '/:id/history', 'driver, admin', "Today's trip location history"],
            ['PUT', '/:id/location', 'driver, admin', 'Body: lat, lng, speed, heading'],
            ['PUT', '/:id/status', 'driver, admin', 'Body: status'],
            ['POST', '/:id/assign', 'driver', 'Assign bus to logged-in driver'],
            ['POST', '/:id/release', 'driver', 'Release bus — idle, isLive false'],
        ])

    add_heading(doc, '7.4 Routes — /api/routes', 2)
    add_table(doc, ['Method', 'Path', 'Notes'],
        [
            ['GET', '/', 'All routes'],
            ['GET', '/:id', 'One route'],
            ['GET', '/:id/buses', 'Live buses on route'],
            ['POST', '/', 'DEPRECATED — use /api/admin/routes'],
            ['PUT', '/:id', 'DEPRECATED'],
            ['DELETE', '/:id', 'Delete route'],
        ])

    add_heading(doc, '7.5 Trips — /api/trips', 2)
    add_table(doc, ['Method', 'Path', 'Auth', 'Body/Notes'],
        [
            ['GET', '/active', 'admin only', 'All ongoing trips'],
            ['GET', '/:id', 'Bearer', 'admin or owning driver'],
            ['POST', '/start', 'driver', '{ busId, routeId? } — creates ongoing trip'],
            ['PUT', '/:id/end', 'driver', 'Completes trip, bus idle'],
        ])

    add_heading(doc, '7.6 Tickets — /api/tickets', 2)
    add_table(doc, ['Method', 'Path', 'Auth', 'Notes'],
        [
            ['POST', '/book', 'passenger', 'Needs ongoing tripId — QR image'],
            ['POST', '/book-simple', 'passenger', 'No trip — JSON QR with ticketId'],
            ['GET', '/my', 'passenger', 'My tickets list'],
            ['POST', '/verify', 'driver', 'Body: ticketId OR qrCode (JSON/raw)'],
            ['PUT', '/:id/cancel', 'passenger', 'Cancel own ticket'],
        ])

    add_heading(doc, '7.7 Stops — /api/stops', 2)
    add_para(doc, 'GET / — active stops. GET /:id — one stop. ?all=1 for all including inactive.')

    add_heading(doc, '7.8 Timetable — /api/timetable', 2)
    add_para(doc, 'GET / — all or ?routeId=. GET /search?from=&to= — search by stop names.')

    add_heading(doc, '7.9 Admin — /api/admin (ALL require admin JWT)', 2)
    add_table(doc, ['Endpoint group', 'Operations'],
        [
            ['/dashboard', 'Stats: live buses, drivers, routes, today trips, passengers, revenue'],
            ['/buses', 'GET list, POST create, PUT update, DELETE'],
            ['/drivers', 'GET list, POST create (password required)'],
            ['/passengers', 'GET list'],
            ['/routes', 'CRUD with stop building from stopId'],
            ['/stops', 'CRUD (delete = soft isActive false)'],
            ['/timetable', 'CRUD'],
            ['/reports/daily', 'Today hourly passengers chart data'],
            ['/reports/revenue', '?days=7 revenue by day'],
        ])

    # 8 SOCKET
    add_heading(doc, '8. Socket.io Real-Time Events', 1)
    add_para(doc, 'Connection URL: same as API (e.g. http://localhost:5000). Transport: websocket.')
    add_para(doc, 'Handshake auth (optional): { token: "<JWT>" }')

    add_heading(doc, '8.1 Client → Server Events', 2)
    add_para(doc, 'passenger:watch — Body: { busId: "..." } — joins room bus:<busId>', bold=True)
    add_para(doc, 'passenger:watchStop — Body: { stopId: "..." } — joins room stop:<stopId>', bold=True)
    add_para(doc, 'driver:location — Driver only. Body:', bold=True)
    add_para(doc, '{ busId, lat, lng, speed, heading, tripId?, timestamp }')
    add_para(doc, 'Server checks: socket.user.role === driver AND bus assigned to this driver.')

    add_heading(doc, '8.2 Server → Client Events', 2)
    add_para(doc, 'buses:locations — Every 3 seconds to ALL clients. Array of bus objects with busId, busName, lat, lng, speed, routeId, eta, seatsAvailable.', bold=True)
    add_para(doc, 'bus:update — To room bus:<busId> on each driver location update.', bold=True)
    add_para(doc, 'stop:eta — To room stop:<stopId> for each route stop when driver moves. { stopId, busId, etaMinutes, distanceKm, updatedAt }', bold=True)

    add_heading(doc, '8.3 Seats Available Calculation', 2)
    add_para(doc, 'seatsAvailable = bus.capacity - ongoing_trip.passengersCount (minimum logic)')

    # 9 AI ETA
    add_heading(doc, '9. AI ETA Microservice (ai-eta/)', 1)
    add_para(doc, 'Run: cd ai-eta && pip install -r requirements.txt && uvicorn main:app --reload --port 8000')
    add_para(doc, 'POST /predict-eta — Request body (snake_case):')
    add_para(doc, 'bus_lat, bus_lng, stop_lat, stop_lng, current_speed, hour_of_day, day_of_week')
    add_para(doc, 'Response: { eta_minutes, distance_km }')
    add_para(doc, 'Rush hour (8-9, 17-18): multiplier 1.4. Night: 0.8. Backend etaService.js calls this when ETA_SERVICE_URL set; else math fallback (haversine + speed).')

    # 10 SEED
    add_heading(doc, '10. Seed Data (npm run seed)', 1)
    add_para(doc, 'WARNING: Seed wipes all collections first (deleteMany).')
    add_para(doc, 'Creates: 3 users (admin, driver, passenger), 8 Botad stops, 3 routes (R1,R2,R3), 1 bus GJ-11-T-2847 assigned to driver on R1, 4 timetable entries for R1.')
    add_para(doc, 'Botad coordinates center: ~22.1647, 71.6661')

    # 11 FRONTEND
    add_heading(doc, '11. React Web Frontend', 1)
    add_heading(doc, '11.1 Routes (App.jsx)', 2)
    add_table(doc, ['URL', 'Page', 'Who'],
        [
            ['/', 'PassengerMap', 'Public — live map, no login'],
            ['/login', 'Login', 'Admin login'],
            ['/admin', 'AdminDashboard', 'Admin only'],
            ['/admin/stops', 'AdminStops', 'CRUD stops'],
            ['/admin/routes', 'AdminRoutes', 'CRUD routes + OSRM polyline'],
            ['/admin/timetable', 'AdminTimetable', 'Schedules'],
            ['/admin/users', 'AdminUsers', 'Drivers + passengers'],
        ])
    add_heading(doc, '11.2 Admin Dashboard', 2)
    add_para(doc, 'Shows: live buses count, passengers today, revenue, routes. Live Google Map (needs VITE_GOOGLE_MAPS_KEY). Hourly passenger bar chart from /api/admin/reports/daily. Fleet table.')
    add_heading(doc, '11.3 Passenger Map', 2)
    add_para(doc, 'Public page at /. Socket se live bus markers. Google Maps. No ticket booking on web.')

    # 12 DRIVER
    add_heading(doc, '12. Flutter Driver App — Complete Guide', 1)
    add_heading(doc, '12.1 Files', 2)
    add_bullet(doc, 'main.dart — LoginScreen start')
    add_bullet(doc, 'constants.dart — API_BASE_URL, MAPTILER_KEY via --dart-define')
    add_bullet(doc, 'services/auth_service.dart — login, token in SharedPreferences')
    add_bullet(doc, 'services/socket_service.dart — connect + emit driver:location')
    add_bullet(doc, 'services/gps_service.dart — Geolocator high accuracy stream')
    add_heading(doc, '12.2 Screen Flow', 2)
    add_para(doc, 'LoginScreen → POST /api/auth/driver/login → HomeScreen')
    add_para(doc, 'HomeScreen: GET /api/buses. If no bus assigned: list buses → POST assign. If assigned: Start Trip OR Scan Ticket.')
    add_para(doc, 'Start Trip: POST /api/trips/start { busId, routeId } → TripScreen with tripId')
    add_para(doc, 'TripScreen: GPS every update + Socket emit every 3 sec. MapTiler map. End Trip: PUT /trips/:id/end + POST /buses/:id/release')
    add_para(doc, 'QRScannerScreen: POST /api/tickets/verify { qrCode }')
    add_heading(doc, '12.3 Run Command', 2)
    add_para(doc, 'flutter run --dart-define=API_BASE_URL=http://YOUR_LAN_IP:5000 --dart-define=MAPTILER_KEY=your_key')
    add_para(doc, 'Android emulator default: http://10.0.2.2:5000')

    # 13 PASSENGER
    add_heading(doc, '13. Flutter Passenger App — Complete Guide', 1)
    add_heading(doc, '13.1 Bottom Navigation (AppShell)', 2)
    add_table(doc, ['Tab', 'Screen', 'Function'],
        [
            ['Home', 'HomeScreen', 'Feature cards — navigate to other screens'],
            ['Find bus', 'FindBusScreen', 'From/To stop → list matching buses'],
            ['Live map', 'MapScreen', 'All live buses on map, destination alert'],
            ['More', 'MoreScreen', 'Login, tickets, timetable, help'],
        ])
    add_heading(doc, '13.2 All Screens', 2)
    screens = [
        ('FindBusScreen', 'Pick from/to stops → GET /api/buses?from=&to='),
        ('BusTrackScreen', 'Track one bus — MapScreen with targetBusId'),
        ('MapScreen', 'Socket buses:locations, stop markers, route polyline from GET /api/routes/:id, arrival alert <500m'),
        ('NearbyStationsScreen', 'Stops sorted by GPS distance'),
        ('TimetableScreen', 'GET /api/routes — show schedules'),
        ('TicketScreen', 'POST /api/tickets/book-simple — show QR (needs login)'),
        ('MyTicketsScreen', 'GET /api/tickets/my'),
        ('LoginScreen / RegisterScreen', 'Auth APIs'),
        ('SearchBusesScreen', 'Alternate bus search'),
        ('BusListScreen', 'GET /api/buses list'),
        ('InfoScreen', 'Static help'),
        ('BusDetailsScreen', 'Per-stop ETA — note: socket watchStop may need backend stop _id'),
    ]
    for name, desc in screens:
        add_bullet(doc, f'{name}: {desc}')
    add_para(doc, 'Passenger can use map WITHOUT login. Login required for ticket booking.')

    # 14 FLOWS
    add_heading(doc, '14. Complete User Flows (Step by Step)', 1)

    add_heading(doc, '14.1 Admin — First Time Setup', 2)
    add_numbered(doc, 'Backend: npm run seed')
    add_numbered(doc, 'Web login: 9999999999 / admin123')
    add_numbered(doc, 'Admin → Stops: add/edit Botad stops with lat/lng')
    add_numbered(doc, 'Admin → Routes: select stops in order, OSRM draws polyline')
    add_numbered(doc, 'Admin → Timetable: departure times per route')
    add_numbered(doc, 'Admin → Users: create drivers with password')
    add_numbered(doc, 'Admin → Buses: assign bus to route and driver')

    add_heading(doc, '14.2 Driver — Daily Trip', 2)
    add_numbered(doc, 'Open driver app → login 9876543210 / driver123')
    add_numbered(doc, 'Select bus if not assigned')
    add_numbered(doc, 'Tap Start Trip → GPS starts → passengers see bus on map')
    add_numbered(doc, 'Passengers board → scan QR → verify success')
    add_numbered(doc, 'End Trip → bus released → trip completed in DB')

    add_heading(doc, '14.3 Passenger — Find & Track Bus', 2)
    add_numbered(doc, 'Open passenger app → Find bus tab')
    add_numbered(doc, 'Select From: Botad Bus Stand, To: Botad College')
    add_numbered(doc, 'See buses on route → tap → live track')
    add_numbered(doc, 'Or Live map tab — see all buses')

    add_heading(doc, '14.4 Passenger — Book Ticket', 2)
    add_numbered(doc, 'More → Login → 9123456789 / pass123')
    add_numbered(doc, 'Ticket screen → select from/to → Get ticket')
    add_numbered(doc, 'QR shows JSON with ticketId')
    add_numbered(doc, 'Show QR to driver → driver scans → verified')

    # 15 ENV
    add_heading(doc, '15. Environment Variables', 1)
    add_heading(doc, '15.1 Backend .env', 2)
    add_table(doc, ['Variable', 'Example', 'Purpose'],
        [
            ['PORT', '5000', 'Server port'],
            ['MONGODB_URI', 'mongodb://127.0.0.1:27017/botad-bus-tracker', 'Database'],
            ['JWT_SECRET', 'long-random-string', 'REQUIRED — token signing'],
            ['JWT_EXPIRE', '7d', 'Token expiry'],
            ['CORS_ORIGINS', 'http://localhost:3000', 'Allowed web origins'],
            ['ETA_SERVICE_URL', 'http://localhost:8000', 'Optional AI ETA'],
        ])
    add_heading(doc, '15.2 Frontend .env', 2)
    add_table(doc, ['Variable', 'Purpose'],
        [['VITE_BACKEND_URL', 'REST API'], ['VITE_SOCKET_URL', 'Socket.io'], ['VITE_GOOGLE_MAPS_KEY', 'Maps']])
    add_heading(doc, '15.3 Flutter dart-define', 2)
    add_table(doc, ['Define', 'Default', 'Purpose'],
        [['API_BASE_URL', 'http://10.0.2.2:5000', 'API + Socket'], ['MAPTILER_KEY', 'YOUR_MAPTILER_KEY', 'Map tiles']])

    # 16 SETUP
    add_heading(doc, '16. Setup & Run — Step by Step (Windows)', 1)
    add_heading(doc, '16.1 Prerequisites', 2)
    add_bullet(doc, 'Node.js 18+, MongoDB (local or Atlas), Python 3.10+ (for ai-eta), Flutter SDK, Git')
    add_heading(doc, '16.2 Backend', 2)
    add_numbered(doc, 'cd botad-bus-tracker/backend')
    add_numbered(doc, 'copy .env.example .env — edit MONGODB_URI, JWT_SECRET')
    add_numbered(doc, 'npm install')
    add_numbered(doc, 'npm run seed')
    add_numbered(doc, 'npm run dev — listen on port 5000')
    add_heading(doc, '16.3 Frontend', 2)
    add_numbered(doc, 'cd frontend')
    add_numbered(doc, 'copy .env.example .env')
    add_numbered(doc, 'npm install && npm run dev')
    add_heading(doc, '16.4 AI ETA (optional)', 2)
    add_numbered(doc, 'cd ai-eta && pip install -r requirements.txt')
    add_numbered(doc, 'uvicorn main:app --reload --port 8000')
    add_heading(doc, '16.5 Flutter Apps', 2)
    add_numbered(doc, 'cd driver-app OR passenger-app')
    add_numbered(doc, 'flutter pub get')
    add_numbered(doc, 'flutter run --dart-define=API_BASE_URL=http://192.168.x.x:5000 --dart-define=MAPTILER_KEY=xxx')
    add_para(doc, 'Physical phone: use PC LAN IP, not localhost.')

    # 17 DEPLOY
    add_heading(doc, '17. Deployment', 1)
    add_table(doc, ['Component', 'Platform', 'Notes'],
        [
            ['Backend', 'Render.com', 'Set env vars, Web Service'],
            ['MongoDB', 'Atlas', 'Connection string in MONGODB_URI'],
            ['Frontend', 'Vercel', 'Build + VITE_* env'],
            ['APK', 'GitHub Actions', 'Secrets: MAPTILER_KEY, vars: API_BASE_URL'],
        ])
    add_para(doc, 'Production: set CORS_ORIGINS to Vercel URL. Strong JWT_SECRET. Never commit .env or API keys.')

    # 18 SECURITY
    add_heading(doc, '18. Security Notes', 1)
    add_bullet(doc, 'Passwords hashed with bcrypt (12 rounds)')
    add_bullet(doc, 'JWT on protected routes')
    add_bullet(doc, 'driver:location requires valid driver token + bus ownership')
    add_bullet(doc, 'Trip history and bus history — driver own bus or admin')
    add_bullet(doc, 'GET /api/trips/active — admin only')
    add_bullet(doc, 'Ticket verify — driver role only')
    add_bullet(doc, 'Auth rate limit 5/min')
    add_bullet(doc, 'MapTiler/Google keys via env only — not in source')

    # 19 ISSUES
    add_heading(doc, '19. Known Issues & Tips', 1)
    add_bullet(doc, 'Route embedded stops from seed may lack _id — stop:eta rooms work best when routes built via admin with stopId')
    add_bullet(doc, 'MongoDB must be running before backend starts')
    add_bullet(doc, 'Emulator: 10.0.2.2 = host machine localhost')
    add_bullet(doc, 'Old tickets before ticketId in QR need re-booking')
    add_bullet(doc, 'Web passenger map does not book tickets — use Flutter app')
    add_bullet(doc, 'legacy_admin in backend/public is deprecated — use React admin')

    # 20 GLOSSARY
    add_heading(doc, '20. Glossary', 1)
    add_table(doc, ['Term', 'Meaning'],
        [
            ['ETA', 'Estimated Time of Arrival — kitni der mein bus aayegi'],
            ['JWT', 'JSON Web Token — login session token'],
            ['Socket.io', 'Real-time bidirectional communication'],
            ['QR Code', 'Ticket verification barcode'],
            ['OSRM', 'Open Source Routing Machine — road polyline'],
            ['Polyline', 'Encoded map route line'],
            ['Seed', 'Demo database data script'],
            ['CRUD', 'Create Read Update Delete'],
            ['CORS', 'Cross-Origin Resource Sharing — browser security'],
        ])

    # APPENDIX A - FILES
    add_heading(doc, 'Appendix A — Har Important File Ka Kaam', 1)
    files_backend = [
        ('server.js', 'Express app, routes mount, Socket.io init, MongoDB connect'),
        ('src/config/db.js', 'mongoose.connect, isDbReady() helper'),
        ('src/middleware/auth.js', 'protect + authorize JWT middleware'),
        ('src/middleware/errorHandler.js', 'Global error JSON response'),
        ('src/utils/seed.js', 'Demo data — users, stops, routes, bus, timetable'),
        ('src/utils/distance.js', 'haversineKm, estimateEtaMinutes (rush hour math)'),
        ('src/utils/generateToken.js', 'JWT sign helper'),
        ('src/services/etaService.js', 'Calls ai-eta OR math fallback'),
        ('src/socket/index.js', 'All Socket.io events + 3s broadcast'),
        ('src/controllers/authController.js', 'register, login, driverLogin, me, refresh'),
        ('src/controllers/busController.js', 'Buses CRUD, nearby, assign, release, GPS'),
        ('src/controllers/tripController.js', 'startTrip, endTrip, getActive, getById'),
        ('src/controllers/ticketController.js', 'book, book-simple, verify, cancel, my'),
        ('src/controllers/adminController.js', 'Dashboard, buses, drivers, reports'),
        ('src/controllers/routeAdminController.js', 'Admin route builder with stops'),
        ('src/controllers/stopController.js', 'Public + admin stops'),
        ('src/controllers/timetableController.js', 'Timetable read + admin CRUD'),
        ('src/controllers/routeController.js', 'Public routes read'),
        ('tests/app.test.js', 'Jest: health, trip start/end, ETA fallback'),
    ]
    add_para(doc, 'Backend files:', bold=True)
    for f, d in files_backend:
        add_bullet(doc, f'{f} — {d}')

    files_fe = [
        ('App.jsx', 'React Router — /, /login, /admin/*'),
        ('pages/PassengerMap.jsx', 'Public Google Map + socket buses'),
        ('pages/Login.jsx', 'Admin login form'),
        ('pages/AdminDashboard.jsx', 'Stats + chart + fleet table'),
        ('pages/admin/AdminStops.jsx', 'Stop CRUD form + table'),
        ('pages/admin/AdminRoutes.jsx', 'Route builder + OSRM polyline fetch'),
        ('pages/admin/AdminTimetable.jsx', 'Timetable management'),
        ('pages/admin/AdminUsers.jsx', 'Drivers + passengers + create driver'),
        ('components/AdminLayout.jsx', 'Sidebar navigation shell'),
        ('hooks/useSocket.js', 'Socket.io buses state hook'),
        ('services/api.js', 'Axios + all API functions'),
    ]
    add_para(doc, 'Frontend files:', bold=True)
    for f, d in files_fe:
        add_bullet(doc, f'{f} — {d}')

    files_driver = [
        ('main.dart', 'App entry → LoginScreen'),
        ('constants.dart', 'API_BASE_URL, MAPTILER_KEY from dart-define'),
        ('screens/login_screen.dart', 'Driver phone/password login'),
        ('screens/home_screen.dart', 'Assign bus, Start Trip, Scan QR buttons'),
        ('screens/trip_screen.dart', 'GPS map + socket emit + End Trip'),
        ('screens/qr_scanner_screen.dart', 'mobile_scanner + verify API'),
        ('services/auth_service.dart', 'Token in SharedPreferences'),
        ('services/socket_service.dart', 'driver:location emit'),
        ('services/gps_service.dart', 'Geolocator permission + stream'),
    ]
    add_para(doc, 'Driver app files:', bold=True)
    for f, d in files_driver:
        add_bullet(doc, f'{f} — {d}')

    files_pass = [
        ('main.dart', 'BotadPassengerApp → AppShell'),
        ('screens/app_shell.dart', 'Bottom nav: Home, Find, Map, More'),
        ('screens/home_screen.dart', 'Feature grid + Find bus banner'),
        ('screens/find_bus_screen.dart', 'From/To stop picker → bus list'),
        ('screens/map_screen.dart', 'Live map all buses + polyline + arrival alert'),
        ('screens/bus_track_screen.dart', 'Single bus tracking wrapper'),
        ('screens/ticket_screen.dart', 'Book ticket + QR display'),
        ('screens/my_tickets_screen.dart', 'Ticket history'),
        ('screens/timetable_screen.dart', 'Route schedules'),
        ('screens/nearby_stations_screen.dart', 'All stops by distance'),
        ('screens/more_screen.dart', 'Account, login, help links'),
        ('screens/login_screen.dart', 'Passenger login'),
        ('screens/register_screen.dart', 'New passenger register'),
        ('services/auth_service.dart', 'Login/register/logout'),
        ('services/socket_service.dart', 'PassengerSocketService — buses:locations'),
        ('services/stops_service.dart', 'Load stops from routes API'),
        ('theme/app_theme.dart', 'Material theme colors/fonts'),
        ('widgets/brand_logo.dart', 'App logo widget'),
        ('widgets/feature_card.dart', 'Home grid cards'),
        ('widgets/stop_search_sheet.dart', 'Stop search bottom sheet'),
    ]
    add_para(doc, 'Passenger app files:', bold=True)
    for f, d in files_pass:
        add_bullet(doc, f'{f} — {d}')

    # APPENDIX B - JSON EXAMPLES
    add_heading(doc, 'Appendix B — Full API JSON Examples', 1)
    add_heading(doc, 'B.1 Register Passenger', 2)
    add_para(doc, 'Request POST /api/auth/register:')
    add_para(doc, '{"name":"Rahul","phone":"9000000001","password":"mypass123","email":"rahul@test.com"}')
    add_para(doc, 'Response 201:')
    add_para(doc, '{"success":true,"token":"eyJhbG...","user":{"id":"...","name":"Rahul","phone":"9000000001","role":"passenger"}}')

    add_heading(doc, 'B.2 Book Simple Ticket', 2)
    add_para(doc, 'Request POST /api/tickets/book-simple (Bearer passenger):')
    add_para(doc, '{"fromStop":"Botad Bus Stand","toStop":"Botad College","price":15}')
    add_para(doc, 'Response — data.qrCode contains JSON string like:')
    add_para(doc, '{"ticketId":"665abc...","passengerId":"...","fromStop":"Botad Bus Stand","toStop":"Botad College","expiresAt":"2026-05-19T10:00:00.000Z"}')

    add_heading(doc, 'B.3 Verify Ticket (Driver Scan)', 2)
    add_para(doc, 'Request POST /api/tickets/verify:')
    add_para(doc, '{"qrCode":"{\\"ticketId\\":\\"665abc...\\",...}"}')
    add_para(doc, 'Success: {"success":true,"message":"Verified","ticket":{...,"status":"used"}}')
    add_para(doc, 'Fail used: {"success":false,"message":"Ticket used"}')
    add_para(doc, 'Fail expired: {"success":false,"message":"QR expired"}')

    add_heading(doc, 'B.4 Socket buses:locations sample', 2)
    add_para(doc, '[{"busId":"665...","busName":"B1","lat":22.1647,"lng":71.6661,"speed":35,"routeNumber":"R1","eta":8,"seatsAvailable":47}]')

    # APPENDIX C - HTTP STATUS
    add_heading(doc, 'Appendix C — HTTP Status Codes Used', 1)
    add_table(doc, ['Code', 'Meaning', 'When'],
        [
            ['200', 'OK', 'Successful GET/PUT'],
            ['201', 'Created', 'Register, book ticket, start trip, create bus'],
            ['400', 'Bad Request', 'Validation fail, trip not ongoing, bus assigned'],
            ['401', 'Unauthorized', 'No/invalid JWT'],
            ['403', 'Forbidden', 'Wrong role, not your bus/ticket'],
            ['404', 'Not Found', 'Bus/trip/ticket not found'],
            ['500', 'Server Error', 'Unhandled exception via errorHandler'],
        ])

    # APPENDIX D - PORTS
    add_heading(doc, 'Appendix D — Ports & URLs Quick Reference', 1)
    add_table(doc, ['Service', 'Local URL', 'Production example'],
        [
            ['Backend API', 'http://localhost:5000', 'https://botad-bus-tracker.onrender.com'],
            ['Health', 'http://localhost:5000/health', 'Same + /health'],
            ['Web', 'http://localhost:3000', 'Vercel URL'],
            ['AI ETA', 'http://localhost:8000', 'Deploy separately'],
            ['MongoDB', 'mongodb://127.0.0.1:27017/botad-bus-tracker', 'Atlas connection string'],
        ])

    # APPENDIX E - TICKET PRICE
    add_heading(doc, 'Appendix E — Business Rules (Chhoti Details)', 1)
    add_bullet(doc, 'Default ticket price: ₹15 (TICKET_PRICE constant in ticketController.js)')
    add_bullet(doc, 'QR validity: 24 hours (QR_EXPIRE_HOURS = 24)')
    add_bullet(doc, 'Ticket status flow: active → used (verify) OR cancelled (expire/cancel)')
    add_bullet(doc, 'Driver location broadcast interval: 3 seconds (Timer in trip_screen + socket broadcast)')
    add_bullet(doc, 'GPS distance filter: 10 meters (GpsService driver app)')
    add_bullet(doc, 'Passenger arrival notification: when bus within 500 meters of selected stop (map_screen.dart)')
    add_bullet(doc, 'Bus filter by route: from stop must come BEFORE to stop in route.stops order')
    add_bullet(doc, 'bcrypt salt rounds: 12')
    add_bullet(doc, 'JWT default expiry: 7 days (JWT_EXPIRE env)')
    add_bullet(doc, 'Auth rate limit: 5 requests per minute on /api/auth/*')

    add_heading(doc, 'Document End', 1)
    add_para(doc, 'Botad Bus Tracker — Full Technical Documentation (Hindi + English). Har module, API, screen, file, aur flow cover kiya gaya hai.')
    add_para(doc, 'File location: botad-bus-tracker/docs/Botad_Bus_Tracker_Complete_Documentation.docx')

    out = r'd:\app\botad bus system\botad-bus-tracker\docs\Botad_Bus_Tracker_Complete_Documentation.docx'
    doc.save(out)
    print(f'Saved: {out}')
    return out

if __name__ == '__main__':
    build()
